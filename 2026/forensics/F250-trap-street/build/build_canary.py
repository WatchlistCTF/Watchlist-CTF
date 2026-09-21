#!/usr/bin/env python3
"""F250 canary builder - decima_subject_register.xlsx"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

CANARY_ID = "NLA-2026-04"
OUT = "decima_subject_register.xlsx"

ROWS = [
 ("Anthony Marconi","878-20-8492","Enforcer","Washington DC","MONITORING","2026-02-01"),
 ("Claire Mahoney","285-42-6165","Activist","New York","FLAGGED","2026-02-18"),
 ("Leon Tao","230-18-2768","Accountant","Los Angeles","CLEARED","2026-03-10"),
 ("Lionel Fusco","497-59-2094","NYPD Detective","New York","MONITORING","2026-01-28"),
 ("Nathan Ingram","200-37-4436","Technology Exec","New York","CLOSED",""),
 ("Hanna Frey","681-58-7511","Student","Los Angeles","CLOSED",""),
 ("Rania Adler","844-35-5420","Journalist","Chicago","MONITORING","2026-01-19"),
 ("Will Ingram","540-11-7718","Physician","New York","CLEARED","2026-01-10"),
 ("Jason Greenfield","303-11-1981","Hacktivist","New York","FLAGGED","2026-01-08"),
 ("Peter Yogorov","898-81-4086","Organized Crime","New York","MONITORING","2026-02-06"),
 ("Clara Holm","334-63-7287","Physician","New York","CLEARED","2026-03-07"),
 ("Joss Carter","604-63-4487","NYPD Detective","New York","CLOSED",""),
 ("Esme Romano","806-48-1321","Financier","New York","CLEARED","2026-02-28"),
 ("Dani Silva","816-92-2643","NYPD Officer","New York","CLEARED","2026-01-13"),
 ("Alicia Corwin","349-37-8234","Government","Washington DC","CLOSED",""),
 ("Denton Weeks","209-88-6389","Government","Washington DC","CLOSED",""),
 ("Harper Rose","276-21-4419","Consultant","New York","CLEARED","2026-02-03"),
 ("Harold Finch","577-57-8425","Software Engineer","New York","FLAGGED","2026-01-20"),
 ("Peter Collier","788-27-7325","Activist","New York","FLAGGED","2026-03-16"),
 ("Omar Haddad","518-39-5087","Engineer","New York","CLEARED","2026-03-05"),
 ("Bianca Marino","844-80-4220","Attorney","Los Angeles","CLEARED","2026-01-24"),
 ("Daniel Casey","280-63-1776","Programmer","New York","CLEARED","2026-02-20"),
 ("John Reese","724-42-4906","Former CIA","Los Angeles","FLAGGED","2026-01-02"),
 ("Sofia Campos","810-72-5807","Student","New York","CLEARED","2026-02-14"),
 ("Joey Durban","329-39-8852","Security","New York","CLEARED","2026-03-03"),
 ("Caleb Phipps","408-12-2133","Student","Washington DC","CLEARED","2026-02-07"),
 ("Sameen Shaw","455-17-1763","Former ISA","New York","FLAGGED","2026-02-15"),
 ("Carl Elias","743-83-3155","Organized Crime","New York","FLAGGED","2026-02-12"),
 ("Zoe Morgan","661-52-9551","Consultant","New York","MONITORING","2026-02-05"),
 ("Mr. Hersh","218-70-6856","Government Operative","Washington DC","MONITORING","2026-03-02"),
 ("Samantha Groves","812-91-2226","Hacker","New York","FLAGGED","2026-01-01"),
 ("Tomas Brandt","518-50-3238","Courier","New York","CLEARED","2026-01-24"),
 ("Kenji Tanaka","759-57-1728","Compliance","New York","CLEARED","2026-01-15"),
 ("Pavel Sokolov","286-97-8754","Organized Crime","New York","MONITORING","2026-02-12"),
 ("Logan Pierce","230-73-1238","Technology Exec","Los Angeles","MONITORING","2026-02-26"),
 ("Grace Hendricks","212-87-2183","Artist","New York","MONITORING","2026-02-19"),
 ("Dmitri Falk","318-42-7819","Arms Dealer","Chicago","FLAGGED","2026-01-21"),
 ("Felix Berg","794-68-8217","Programmer","New York","CLEARED","2026-02-24"),
 ("Marcus Chen","731-75-1488","Field Engineer","Washington DC","CLEARED","2026-03-03"),
 ("Nadia Okafor","692-12-4771","Researcher","New York","MONITORING","2026-03-03"),
 ("Julia Voss","829-94-8967","Logistics","Washington DC","MONITORING","2026-02-25"),
]

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Subject Register"
title=Font(bold=True,size=14,color="1F2A44"); sub=Font(italic=True,size=9,color="888888")
hdr_f=Font(bold=True,color="FFFFFF"); hdr_fl=PatternFill("solid",fgColor="1F2A44")
thin=Side(style="thin",color="CCCCCC"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
sfill={"FLAGGED":PatternFill("solid",fgColor="FCE4E4"),"MONITORING":PatternFill("solid",fgColor="FFF6E0"),
       "CLEARED":PatternFill("solid",fgColor="E8F2E8"),"CLOSED":PatternFill("solid",fgColor="ECECEC")}

ws["A1"]="DECIMA TECHNOLOGIES // NORTHERN LIGHTS - SUBJECT REGISTER (Q4 RELEVANCE REVIEW)"; ws["A1"].font=title
ws["A2"]="CONFIDENTIAL // RELEVANCE CLASSIFICATION // STAGED FOR TRANSFER - DO NOT REDISTRIBUTE"; ws["A2"].font=sub
cols=["Subject","Number","Classification","Region","Status","Last Assessed"]
for i,h in enumerate(cols,start=1):
    c=ws.cell(row=3,column=i,value=h); c.font=hdr_f; c.fill=hdr_fl; c.border=border
    c.alignment=Alignment(horizontal="center")
for r,rec in enumerate(ROWS,start=4):
    for i,v in enumerate(rec,start=1):
        c=ws.cell(row=r,column=i,value=v); c.border=border
    sc=ws.cell(row=r,column=5); sc.fill=sfill.get(rec[4],PatternFill()); sc.alignment=Alignment(horizontal="center")
for col,w in zip("ABCDEF",(18,13,22,15,13,14)): ws.column_dimensions[col].width=w
ws.freeze_panes="A4"

white=Font(color="FFFFFF")
for cell,val in {"A60":"transit","A61":".northernlights","A62":".gg/checkin","A63":f"?id={CANARY_ID}"}.items():
    ws[cell]=val; ws[cell].font=white
ws["A65"]="=A60 & A61 & A62 & A63"; ws["A65"].font=white

ws2=wb.create_sheet("Notes"); ws2["A1"]="Relevance Classification - Notes"; ws2["A1"].font=Font(bold=True)
ws2["A3"]=("Classification reflects current relevance assessment. FLAGGED denotes an active threat "
           "assessment; MONITORING denotes passive observation; CLEARED denotes no current relevance; "
           "CLOSED denotes a terminated record. Reassessment occurs quarterly."); ws2.column_dimensions["A"].width=100
ws3=wb.create_sheet("Approvals"); ws3["A1"]="Approvals"; ws3["A1"].font=Font(bold=True)
ws3["A3"]="Reviewed by:"; ws3["B3"]="Operations Directorate"; ws3["A4"]="Date:"; ws3["B4"]="2026-04-22"
ws3["A5"]="Control ID:"; ws3["B5"]="NL-SR-2026-Q4-0007"
for col,w in zip("AB",(14,30)): ws3.column_dimensions[col].width=w
wb.save(OUT); print("built",OUT,"with",len(ROWS),"subjects")
